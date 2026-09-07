import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))

DOCUMENTS = [
    {
        "md": os.path.join(BASE_DIR, 'legal_technical_report', 'operation_phantom_swipe_report.md'),
        "pdf": os.path.join(BASE_DIR, 'legal_technical_report', 'operation_phantom_swipe_report.pdf'),
        "docx": os.path.join(BASE_DIR, 'legal_technical_report', 'operation_phantom_swipe_report.docx'),
        "title": "Legal-Technical Investigation Report"
    },
    {
        "md": os.path.join(BASE_DIR, 'chain_of_custody', 'chain_of_custody_form.md'),
        "pdf": os.path.join(BASE_DIR, 'chain_of_custody', 'chain_of_custody_form.pdf'),
        "docx": os.path.join(BASE_DIR, 'chain_of_custody', 'chain_of_custody_form.docx'),
        "title": "Chain of Custody Form"
    }
]

def clean_md_inline(text):
    text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    parts = text.split('**')
    res = []
    for idx, part in enumerate(parts):
        if idx % 2 == 1:
            res.append(f"<b>{part}</b>")
        else:
            res.append(part)
    text = "".join(res)
    
    parts = text.split('`')
    res = []
    for idx, part in enumerate(parts):
        if idx % 2 == 1:
            res.append(f"<font name='Courier'>{part}</font>")
        else:
            res.append(part)
    text = "".join(res)
    return text

def convert_md_to_pdf(md_path, pdf_path):
    print(f"[*] Generating PDF: {pdf_path}")
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    NAVY = colors.HexColor('#1A365D')
    BLUE = colors.HexColor('#2B6CB0')
    DARK = colors.HexColor('#2D3748')
    LIGHT_BG = colors.HexColor('#F7FAFC')
    LINE_COLOR = colors.HexColor('#E2E8F0')
    
    title_style = ParagraphStyle('DocTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=NAVY, alignment=TA_CENTER, spaceAfter=8)
    subtitle_style = ParagraphStyle('DocSubTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=12, leading=15, textColor=BLUE, alignment=TA_CENTER, spaceAfter=12)
    meta_style = ParagraphStyle('DocMeta', parent=styles['Normal'], fontName='Helvetica-Oblique', fontSize=9, leading=13, textColor=DARK, alignment=TA_CENTER, spaceAfter=16)
    h1_style = ParagraphStyle('Header1', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=13, leading=17, textColor=NAVY, spaceBefore=14, spaceAfter=6, keepWithNext=True)
    h2_style = ParagraphStyle('Header2', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10.5, leading=14, textColor=BLUE, spaceBefore=10, spaceAfter=4, keepWithNext=True)
    body_style = ParagraphStyle('BodyText', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=13.5, textColor=DARK, alignment=TA_JUSTIFY, spaceAfter=6)
    bullet_style = ParagraphStyle('BulletText', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=13, textColor=DARK, leftIndent=15, spaceAfter=4)
    table_cell = ParagraphStyle('TableCell', parent=styles['Normal'], fontName='Helvetica', fontSize=8, leading=10.5, textColor=DARK)
    table_header = ParagraphStyle('TableHeader', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8.5, leading=11, textColor=colors.white)

    story = []
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_rows = []
    
    for line in lines:
        line_str = line.strip()
        if line_str.startswith('|'):
            in_table = True
            parts = [p.strip() for p in line_str.split('|')[1:-1]]
            if parts and all(c in '-: ' for c in parts[0]):
                continue
            table_rows.append(parts)
            continue
        else:
            if in_table and table_rows:
                col_count = max(len(r) for r in table_rows)
                formatted_rows = []
                for r_idx, row in enumerate(table_rows):
                    formatted_row = []
                    for c_idx, cell_text in enumerate(row):
                        clean_cell = clean_md_inline(cell_text)
                        if r_idx == 0:
                            formatted_row.append(Paragraph(clean_cell, table_header))
                        else:
                            formatted_row.append(Paragraph(clean_cell, table_cell))
                    while len(formatted_row) < col_count:
                        formatted_row.append(Paragraph('', table_cell))
                    formatted_rows.append(formatted_row)
                
                t = Table(formatted_rows, repeatRows=1)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), NAVY),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('GRID', (0, 0), (-1, -1), 0.5, LINE_COLOR),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, LIGHT_BG]),
                    ('TOPPADDING', (0, 0), (-1, -1), 4),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                    ('LEFTPADDING', (0, 0), (-1, -1), 4),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                ]))
                story.append(t)
                story.append(Spacer(1, 8))
                table_rows = []
                in_table = False

        if not line_str:
            story.append(Spacer(1, 4))
            continue
            
        if line_str.startswith('# '):
            clean = clean_md_inline(line_str[2:])
            story.append(Paragraph(clean, title_style))
        elif line_str.startswith('## '):
            clean = clean_md_inline(line_str[3:])
            story.append(Paragraph(clean, subtitle_style))
        elif line_str.startswith('### '):
            clean = clean_md_inline(line_str[4:])
            story.append(Paragraph(clean, h1_style))
            story.append(HRFlowable(width="100%", thickness=1, color=NAVY, spaceAfter=6, spaceBefore=2))
        elif line_str.startswith('#### '):
            clean = clean_md_inline(line_str[5:])
            story.append(Paragraph(clean, h2_style))
        elif line_str.startswith('- ') or line_str.startswith('* '):
            clean = clean_md_inline(line_str[2:])
            story.append(Paragraph(f"• {clean}", bullet_style))
        elif line_str.startswith('1. ') or line_str.startswith('2. ') or line_str.startswith('3. ') or line_str.startswith('4. ') or line_str.startswith('5. '):
            clean = clean_md_inline(line_str)
            story.append(Paragraph(clean, bullet_style))
        elif line_str.startswith('```'):
            continue
        elif line_str.startswith('┌') or line_str.startswith('│') or line_str.startswith('└') or line_str.startswith('├') or line_str.startswith('▼') or line_str.startswith('▲'):
            clean_diag = line_str.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace(' ', '&nbsp;')
            story.append(Paragraph(f"<font name='Courier' size=7>{clean_diag}</font>", body_style))
        else:
            clean = clean_md_inline(line_str)
            if 'Case File Reference:' in line_str or 'Case Name:' in line_str or 'Date of Issue:' in line_str:
                story.append(Paragraph(clean, meta_style))
            else:
                story.append(Paragraph(clean, body_style))
                
    doc.build(story)
    print(f"  [+] PDF created: {pdf_path}")

def convert_md_to_docx(md_path, docx_path):
    print(f"[*] Generating DOCX: {docx_path}")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_table = False
    table_rows = []
    
    for line in lines:
        line_str = line.strip()
        if line_str.startswith('|'):
            in_table = True
            parts = [p.strip() for p in line_str.split('|')[1:-1]]
            if parts and all(c in '-: ' for c in parts[0]):
                continue
            table_rows.append(parts)
            continue
        else:
            if in_table and table_rows:
                col_count = max(len(r) for r in table_rows)
                t = doc.add_table(rows=len(table_rows), cols=col_count)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                for r_idx, row_data in enumerate(table_rows):
                    row = t.rows[r_idx]
                    for c_idx, cell_value in enumerate(row_data):
                        cell = row.cells[c_idx]
                        clean_cell = cell_value.replace('**', '').replace('`', '')
                        cell.text = clean_cell
                        
                        if r_idx == 0:
                            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1A365D"/>')
                            cell._tc.get_or_add_tcPr().append(shd)
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    run.font.bold = True
                                    run.font.color.rgb = RGBColor(255, 255, 255)
                                    run.font.size = Pt(8.5)
                        else:
                            if r_idx % 2 == 1:
                                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F7FAFC"/>')
                                cell._tc.get_or_add_tcPr().append(shd)
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    run.font.size = Pt(8.5)
                                    run.font.color.rgb = RGBColor(45, 55, 72)
                
                doc.add_paragraph()
                table_rows = []
                in_table = False

        if not line_str:
            continue
            
        if line_str.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line_str[2:].replace('**', ''))
            run.font.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(26, 54, 93)
        elif line_str.startswith('## '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line_str[3:].replace('**', ''))
            run.font.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(43, 108, 176)
        elif line_str.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(line_str[4:].replace('**', ''))
            run.font.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(26, 54, 93)
        elif line_str.startswith('#### '):
            p = doc.add_paragraph()
            run = p.add_run(line_str[5:].replace('**', ''))
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(43, 108, 176)
        elif line_str.startswith('- ') or line_str.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(line_str[2:].replace('**', '').replace('`', ''))
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(45, 55, 72)
        else:
            p = doc.add_paragraph()
            clean = line_str.replace('**', '').replace('`', '')
            run = p.add_run(clean)
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(45, 55, 72)

    doc.save(docx_path)
    print(f"  [+] DOCX created: {docx_path}")

def build_all_documents():
    for item in DOCUMENTS:
        print(f"\nProcessing Document: {item['title']}")
        convert_md_to_pdf(item['md'], item['pdf'])
        convert_md_to_docx(item['md'], item['docx'])

if __name__ == '__main__':
    build_all_documents()
