import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
PRES_MD = os.path.join(BASE_DIR, 'presentation', 'presentation_guide.md')
PDF_OUT = os.path.join(BASE_DIR, 'presentation', 'presentation_slides.pdf')
DOCX_OUT = os.path.join(BASE_DIR, 'presentation', 'presentation_slides.docx')

def clean_inline(text):
    text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    parts = text.split('**')
    res = []
    for idx, part in enumerate(parts):
        if idx % 2 == 1:
            res.append(f"<b>{part}</b>")
        else:
            res.append(part)
    text = "".join(res)
    
    parts = text.split('`')
    res = []
    for idx, part in enumerate(parts):
        if idx % 2 == 1:
            res.append(f"<font name='Courier'>{part}</font>")
        else:
            res.append(part)
    text = "".join(res)
    return text

def build_pdf_presentation():
    print(f"[*] Generating Presentation PDF: {PDF_OUT}")
    doc = SimpleDocTemplate(
        PDF_OUT,
        pagesize=landscape(letter),
        rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    NAVY = colors.HexColor('#1A365D')
    BLUE = colors.HexColor('#2B6CB0')
    DARK = colors.HexColor('#2D3748')
    LIGHT_BG = colors.HexColor('#F7FAFC')
    LINE_COLOR = colors.HexColor('#E2E8F0')
    
    title_style = ParagraphStyle('PresTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=20, leading=24, textColor=NAVY, alignment=TA_CENTER, spaceAfter=8)
    subtitle_style = ParagraphStyle('PresSubTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=12, leading=16, textColor=BLUE, alignment=TA_CENTER, spaceAfter=14)
    h1_style = ParagraphStyle('PresH1', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=14, leading=18, textColor=NAVY, spaceBefore=12, spaceAfter=6, keepWithNext=True)
    h2_style = ParagraphStyle('PresH2', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=11, leading=14, textColor=BLUE, spaceBefore=8, spaceAfter=4, keepWithNext=True)
    body_style = ParagraphStyle('PresBody', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=13.5, textColor=DARK, alignment=TA_JUSTIFY, spaceAfter=6)
    bullet_style = ParagraphStyle('PresBullet', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=13, textColor=DARK, leftIndent=15, spaceAfter=4)
    code_style = ParagraphStyle('PresCode', parent=styles['Normal'], fontName='Courier', fontSize=8, leading=11, textColor=NAVY, backColor=LIGHT_BG, borderPadding=4, spaceAfter=6)

    story = []
    with open(PRES_MD, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code = False
    code_block = []

    for line in lines:
        line_str = line.strip()
        
        if line_str.startswith('```'):
            if in_code:
                code_text = "<br/>".join(code_block)
                story.append(Paragraph(code_text, code_style))
                code_block = []
                in_code = False
            else:
                in_code = True
            continue
            
        if in_code:
            code_block.append(clean_inline(line_str).replace(' ', '&nbsp;'))
            continue

        if not line_str:
            story.append(Spacer(1, 4))
            continue
            
        if line_str.startswith('# '):
            clean = clean_inline(line_str[2:])
            story.append(Paragraph(clean, title_style))
        elif line_str.startswith('## '):
            clean = clean_inline(line_str[3:])
            story.append(Paragraph(clean, subtitle_style))
        elif line_str.startswith('### '):
            clean = clean_inline(line_str[4:])
            if "Slide " in line_str:
                story.append(PageBreak())
            story.append(Paragraph(clean, h1_style))
            story.append(HRFlowable(width="100%", thickness=1, color=NAVY, spaceAfter=6, spaceBefore=2))
        elif line_str.startswith('#### '):
            clean = clean_inline(line_str[5:])
            story.append(Paragraph(clean, h2_style))
        elif line_str.startswith('- ') or line_str.startswith('* '):
            clean = clean_inline(line_str[2:])
            story.append(Paragraph(f"• {clean}", bullet_style))
        elif line_str.startswith('1. ') or line_str.startswith('2. ') or line_str.startswith('3. ') or line_str.startswith('4. ') or line_str.startswith('5. '):
            clean = clean_inline(line_str)
            story.append(Paragraph(clean, bullet_style))
        elif line_str.startswith('>'):
            clean = clean_inline(line_str[1:].strip())
            story.append(Paragraph(f"<i>Say to Professor:</i> {clean}", body_style))
        else:
            clean = clean_inline(line_str)
            story.append(Paragraph(clean, body_style))

    doc.build(story)
    print(f"  [+] Presentation PDF created: {PDF_OUT}")

def build_docx_presentation():
    print(f"[*] Generating Presentation DOCX: {DOCX_OUT}")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    with open(PRES_MD, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line_str = line.strip()
        if not line_str or line_str.startswith('```'):
            continue
            
        if line_str.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line_str[2:].replace('**', ''))
            run.font.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(26, 54, 93)
        elif line_str.startswith('## '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line_str[3:].replace('**', ''))
            run.font.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(43, 108, 176)
        elif line_str.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(line_str[4:].replace('**', ''))
            run.font.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(26, 54, 93)
        elif line_str.startswith('- ') or line_str.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(line_str[2:].replace('**', '').replace('`', ''))
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(45, 55, 72)
        else:
            p = doc.add_paragraph()
            clean = line_str.replace('**', '').replace('`', '')
            run = p.add_run(clean)
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(45, 55, 72)

    doc.save(DOCX_OUT)
    print(f"  [+] Presentation DOCX created: {DOCX_OUT}")

if __name__ == '__main__':
    build_pdf_presentation()
    build_docx_presentation()
